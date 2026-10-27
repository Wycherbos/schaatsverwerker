"""
Schaats-uitslagen verwerker
Start met:  pip install streamlit pandas openpyxl
            streamlit run schaats_verwerker.py
"""
import io

import pandas as pd
import streamlit as st
from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell

# ----------------------------------------------------------------------
# CONFIGURATIE
# map: CSV-kolomletter -> Excel-kolomletter
# start_row: eerste rij waar data komt (AANNAME: kopregel staat op rij 8)
# ----------------------------------------------------------------------
TAB_KLASSEMENT = "Klassement"
TAB_AFSTAND = "Afstand"
START_RIJ_KLASSEMENT = 9
START_RIJ_AFSTAND = 9
CSV_HEEFT_KOPREGEL = False  # True als de CSV een titelregel bovenaan heeft

CONFIG = {
    "Sprint": {
        "klassement_map": {"A": "A", "B": "B", "C": "C",
                           "D": "F", "E": "H", "F": "J", "G": "M"},
        "afstand_map": {"A": "A", "B": "B", "C": "C",
                        "D": "H",   # tijd
                        "E": "F",   # baan
                        "G": "J"},  # punten
        "labels_klassement": {"F8": "1e 500m", "H8": "1e 1000m", "J8": "2e 500m"},
        "label_afstand": "1e 1000m",
    },
    "Allround": {
        "klassement_map": {"A": "A", "B": "B", "C": "C",
                           "D": "F", "E": "H", "F": "J", "G": "M"},
        "afstand_map": {"A": "A", "B": "B", "C": "C",
                        "D": "H",   # tijd
                        "E": "J"},  # punten (baan blijft leeg)
        "labels_klassement": {"F8": "500m", "H8": None, "J8": "1500m"},  # H8 volgt de keuze 3000m/5000m
        "label_afstand": None,
    },
}


# ----------------------------------------------------------------------
# Hulpfuncties
# ----------------------------------------------------------------------
def letter_naar_index(letter: str) -> int:
    return ord(letter.upper()) - ord("A")


def lees_csv(bestand) -> pd.DataFrame:
    """Leest CSV met automatische detectie van ; of , als scheidingsteken."""
    bestand.seek(0)
    return pd.read_csv(
        bestand, sep=None, engine="python",
        header=0 if CSV_HEEFT_KOPREGEL else None,
        dtype=str, encoding="utf-8-sig", keep_default_na=False,
    )


def naar_waarde(tekst: str):
    tekst = tekst.strip()
    if tekst == "":
        return None
    try:
        return int(tekst)
    except ValueError:
        pass
    try:
        return float(tekst.replace(",", "."))
    except ValueError:
        return tekst  # tijden zoals '1:12.45' blijven tekst


def zet_cel(ws, adres: str, waarde):
    """Schrijft naar een cel; bij samengevoegde cellen naar de linkerbovenhoek."""
    cel = ws[adres]
    if isinstance(cel, MergedCell):
        for bereik in ws.merged_cells.ranges:
            if adres in bereik:
                cel = ws.cell(row=bereik.min_row, column=bereik.min_col)
                break
    cel.value = waarde


def schrijf_data(ws, df, kolom_map, start_rij):
    for i, rij in enumerate(df.itertuples(index=False)):
        for csv_kol, xl_kol in kolom_map.items():
            idx = letter_naar_index(csv_kol)
            if idx < len(rij):
                zet_cel(ws, f"{xl_kol}{start_rij + i}", naar_waarde(rij[idx]))


def verwerk(excel_bytes, csv_klassement, csv_afstand, soort, afstand_allround):
    cfg = CONFIG[soort]
    wb = load_workbook(io.BytesIO(excel_bytes))

    for tab in (TAB_KLASSEMENT, TAB_AFSTAND):
        if tab not in wb.sheetnames:
            raise ValueError(f"Tabblad '{tab}' niet gevonden. Beschikbaar: {', '.join(wb.sheetnames)}")

    ws_k, ws_a = wb[TAB_KLASSEMENT], wb[TAB_AFSTAND]

    # Koppen / afstandsnamen
    labels = dict(cfg["labels_klassement"])
    label_afstand = cfg["label_afstand"]
    if soort == "Allround":
        labels["H8"] = afstand_allround
        label_afstand = afstand_allround
    for adres, tekst in labels.items():
        zet_cel(ws_k, adres, tekst)
    zet_cel(ws_a, "A6", label_afstand)

    # Data
    schrijf_data(ws_k, lees_csv(csv_klassement), cfg["klassement_map"], START_RIJ_KLASSEMENT)
    schrijf_data(ws_a, lees_csv(csv_afstand), cfg["afstand_map"], START_RIJ_AFSTAND)

    uit = io.BytesIO()
    wb.save(uit)
    return uit.getvalue()


# ----------------------------------------------------------------------
# Gebruikersinterface
# ----------------------------------------------------------------------
st.set_page_config(page_title="Schaats-uitslagen verwerker", page_icon="⛸️")
st.title("⛸️ Schaats-uitslagen verwerker")

soort = st.radio("Type wedstrijd", ["Allround", "Sprint"], horizontal=True)

afstand_allround = None
if soort == "Allround":
    afstand_allround = st.radio("Afstand (2e afstand)", ["3000m", "5000m"], horizontal=True)
    lbl_klassement, lbl_afstand = "Klassement Allround.CSV", f"{afstand_allround}.CSV"
else:
    lbl_klassement, lbl_afstand = "Klassement Sprint.CSV", "1e 1000m.CSV"

excel_bestand = st.file_uploader("Excelbestand (sjabloon)", type=["xlsx", "xlsm"])
csv_1 = st.file_uploader(lbl_klassement, type=["csv"], key=f"k_{soort}")
csv_2 = st.file_uploader(lbl_afstand, type=["csv"], key=f"a_{soort}_{afstand_allround}")

if excel_bestand and csv_1 and csv_2:
    if st.button("Verwerk bestanden", type="primary"):
        try:
            resultaat = verwerk(excel_bestand.getvalue(), csv_1, csv_2, soort, afstand_allround)
            st.success("Klaar! Je kunt het bestand downloaden.")
            st.download_button(
                "⬇️ Download Excelbestand", data=resultaat,
                file_name=f"Uitslag_{soort}.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        except Exception as fout:
            st.error(f"Er ging iets mis: {fout}")
else:
    st.info("Upload het Excelbestand en beide CSV-bestanden om door te gaan.")
